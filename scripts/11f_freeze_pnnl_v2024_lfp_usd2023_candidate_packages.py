from __future__ import annotations

import hashlib
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


SCRIPT_VERSION = "v7.1-pnnl-v2024-lfp-usd2023-candidate-package-2026-08-22"

EXPECTED_SHA256 = (
    "ef50e87012361e2a10d3df97cb61b93d38eae9255539ca76061e0c5feef501e9"
)

TECHNOLOGY = "Lithium-ion LFP"
SOURCE_YEAR = 2023
ESTIMATE_TYPE = "Point"
POWER_SCALES_MW = (1.0, 10.0)
DURATIONS_HR = (4.0, 6.0, 8.0, 10.0)

DISCOUNT_RATE = 0.05
ANALYSIS_HORIZON_YR = 20

EXPECTED_ANCHORS = {
    1.0: {
        "capex_C_E_usd_per_kwh": 402.501,
        "capex_C_P_usd_per_kw": 232.648,
        "fom_E_usd_per_kwh_year": 0.857,
        "fom_P_usd_per_kw_year": 1.951,
        "crep_window_mean_usd_per_kwh": 214.375,
    },
    10.0: {
        "capex_C_E_usd_per_kwh": 362.1665,
        "capex_C_P_usd_per_kw": 173.127,
        "fom_E_usd_per_kwh_year": 0.8795,
        "fom_P_usd_per_kw_year": 1.081,
        "crep_window_mean_usd_per_kwh": 191.0775,
    },
}

REQUIRED_PARAMETERS = (
    "Total Installed Cost ($/kW)",
    "Fixed O&M ($/kW-year)",
    "DC Storage Block ($/kWh)",
)


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def crf(rate: float, years: int) -> float:
    return rate * (1.0 + rate) ** years / ((1.0 + rate) ** years - 1.0)


def affine_fit(
    x: np.ndarray,
    y: np.ndarray,
) -> tuple[float, float, float, float]:
    """
    Fit y = intercept + slope*x.

    Returns:
        intercept, slope, r2, max_abs_pct_error
    """
    x = np.asarray(x, dtype=float)
    y = np.asarray(y, dtype=float)

    X = np.column_stack([np.ones_like(x), x])
    beta, *_ = np.linalg.lstsq(X, y, rcond=None)

    intercept = float(beta[0])
    slope = float(beta[1])

    y_hat = X @ beta
    residual = y - y_hat

    ss_res = float(np.sum(residual**2))
    ss_tot = float(np.sum((y - np.mean(y)) ** 2))
    r2 = 1.0 if ss_tot == 0 else 1.0 - ss_res / ss_tot

    with np.errstate(divide="ignore", invalid="ignore"):
        pct_error = np.where(np.abs(y) > 0, residual / y * 100.0, np.nan)

    max_abs_pct_error = float(np.nanmax(np.abs(pct_error)))

    return intercept, slope, r2, max_abs_pct_error


def assert_close(
    observed: float,
    expected: float,
    label: str,
    tol: float = 1e-9,
) -> None:
    if not np.isclose(observed, expected, atol=tol, rtol=0.0):
        raise RuntimeError(
            f"Regression anchor mismatch for {label}: "
            f"observed={observed}, expected={expected}"
        )


def main() -> None:
    print(f"Script version: {SCRIPT_VERSION}")
    print("NTUST v7.1 PNNL v2024 LFP USD-2023 Candidate Cost-Package Freeze")
    print(f"Script version: {SCRIPT_VERSION}")
    print()

    root = Path(__file__).resolve().parents[1]
    workbook_path = (
        root / "data" / "reference" / "ESGC_Cost_Performance_Database_v2024.xlsx"
    )
    output_dir = root / "results" / "parameter_audit"
    output_dir.mkdir(parents=True, exist_ok=True)

    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    workbook_hash = sha256_file(workbook_path)

    print("Workbook provenance")
    print(f"- Workbook: {workbook_path}")
    print(f"- SHA-256: {workbook_hash}")

    if workbook_hash != EXPECTED_SHA256:
        raise RuntimeError(
            "Workbook SHA-256 does not match the audited 11b-11e workbook.\n"
            f"Expected: {EXPECTED_SHA256}\n"
            f"Observed: {workbook_hash}"
        )

    wb = load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
    )

    if "Database" not in wb.sheetnames:
        raise RuntimeError("Worksheet 'Database' not found.")

    ws = wb["Database"]

    expected_headers = [
        "Technology",
        "Year",
        "Power_MW",
        "Duration_hr",
        "Estimate_type",
        "Parameter_category",
        "Parameter",
        "Value",
    ]

    observed_headers = [
        ws.cell(row=1, column=i).value
        for i in range(1, 9)
    ]

    if observed_headers != expected_headers:
        raise RuntimeError(
            "Database headers do not match audited schema.\n"
            f"Expected: {expected_headers}\n"
            f"Observed: {observed_headers}"
        )

    records: list[dict] = []

    for excel_row, row in enumerate(
        ws.iter_rows(min_row=2, max_col=8, values_only=True),
        start=2,
    ):
        (
            technology,
            year,
            power_mw,
            duration_hr,
            estimate_type,
            parameter_category,
            parameter,
            value,
        ) = row

        if technology != TECHNOLOGY:
            continue

        try:
            year_num = int(year)
            power_num = float(power_mw)
            duration_num = float(duration_hr)
            value_num = float(value)
        except (TypeError, ValueError):
            continue

        if year_num != SOURCE_YEAR:
            continue
        if estimate_type != ESTIMATE_TYPE:
            continue
        if power_num not in POWER_SCALES_MW:
            continue
        if duration_num not in DURATIONS_HR:
            continue
        if parameter not in REQUIRED_PARAMETERS:
            continue

        records.append(
            {
                "excel_row": excel_row,
                "power_mw": power_num,
                "duration_hr": duration_num,
                "parameter_category": parameter_category,
                "parameter": parameter,
                "value": value_num,
                "source_value_cell": f"H{excel_row}",
            }
        )

    raw = pd.DataFrame(records)

    if raw.empty:
        raise RuntimeError("No required source rows extracted.")

    duplicate_check = (
        raw.groupby(["power_mw", "duration_hr", "parameter"])
        .size()
        .reset_index(name="n")
    )
    bad = duplicate_check[duplicate_check["n"] != 1]

    if not bad.empty:
        raise RuntimeError(
            "Expected exactly one source row per power/duration/parameter:\n"
            + bad.to_string(index=False)
        )

    packages: list[dict] = []
    source_rows_output: list[dict] = []

    derived_crf = crf(DISCOUNT_RATE, ANALYSIS_HORIZON_YR)

    print()
    print("Method freeze")
    print("- Cost vintage: PNNL v2024 representing 2023 values")
    print("- Estimate type: Point")
    print("- Cost-linearization window: 4, 6, 8, 10 h")
    print("- 1 MW / 10 MW: retained as separate cost-scale brackets")
    print("- CAPEX form: C_E * E_nameplate + C_P * P_rating")
    print("- FOM form: FOM_E * E_nameplate + FOM_P * P_rating")
    print(
        "- C_rep component basis: DC Storage Block ($/kWh), "
        "window-average over 4/6/8/10 h"
    )
    print("- C_rep window averaging is thesis-specific and explicitly traceable.")
    print("- USD->NTD conversion: NOT performed here")
    print("- CRF application: NOT performed here")
    print(f"- CRF diagnostic from r=5%, n=20 yr: {derived_crf:.10f}")

    for power in POWER_SCALES_MW:
        sub = raw[raw["power_mw"] == power].copy()

        pivot = (
            sub.pivot(
                index="duration_hr",
                columns="parameter",
                values="value",
            )
            .reindex(DURATIONS_HR)
        )

        if pivot.isna().any().any():
            raise RuntimeError(
                f"Missing required values for {power:g} MW:\n{pivot}"
            )

        durations = np.asarray(DURATIONS_HR, dtype=float)

        # CAPEX:
        # TIC ($/kW) = C_P ($/kW) + C_E ($/kWh) * duration (h)
        capex_intercept, capex_slope, capex_r2, capex_max_err = affine_fit(
            durations,
            pivot["Total Installed Cost ($/kW)"].to_numpy(dtype=float),
        )

        capex_C_P = capex_intercept
        capex_C_E = capex_slope

        # FOM:
        # Fixed O&M ($/kW-year) =
        # FOM_P ($/kW-year) + FOM_E ($/kWh-year) * duration (h)
        fom_intercept, fom_slope, fom_r2, fom_max_err = affine_fit(
            durations,
            pivot["Fixed O&M ($/kW-year)"].to_numpy(dtype=float),
        )

        fom_P = fom_intercept
        fom_E = fom_slope

        # Replacement-cost economic basis:
        # PNNL reports DC Storage Block directly in $/kWh.
        # Use the transparent arithmetic mean over the same audited 4-10 h window.
        crep_values = pivot["DC Storage Block ($/kWh)"].to_numpy(dtype=float)

        crep_mean = float(np.mean(crep_values))
        crep_median = float(np.median(crep_values))
        crep_min = float(np.min(crep_values))
        crep_max = float(np.max(crep_values))
        crep_range_pct_of_mean = (
            (crep_max - crep_min) / crep_mean * 100.0
        )

        # Diagnostic only: affine fit to total DC-storage-block cost per kW.
        # Total block cost per kW = reported $/kWh * duration.
        sb_total_per_kw = crep_values * durations
        (
            sb_power_intercept,
            sb_energy_slope,
            sb_r2,
            sb_max_err,
        ) = affine_fit(durations, sb_total_per_kw)

        anchor = EXPECTED_ANCHORS[power]

        assert_close(
            capex_C_E,
            anchor["capex_C_E_usd_per_kwh"],
            f"{power:g} MW CAPEX C_E",
            tol=1e-6,
        )
        assert_close(
            capex_C_P,
            anchor["capex_C_P_usd_per_kw"],
            f"{power:g} MW CAPEX C_P",
            tol=1e-6,
        )
        assert_close(
            fom_E,
            anchor["fom_E_usd_per_kwh_year"],
            f"{power:g} MW FOM_E",
            tol=1e-6,
        )
        assert_close(
            fom_P,
            anchor["fom_P_usd_per_kw_year"],
            f"{power:g} MW FOM_P",
            tol=1e-6,
        )
        assert_close(
            crep_mean,
            anchor["crep_window_mean_usd_per_kwh"],
            f"{power:g} MW C_rep window mean",
            tol=1e-9,
        )

        source_rows = (
            sub.sort_values(["parameter", "duration_hr"])
            [["excel_row", "duration_hr", "parameter", "source_value_cell"]]
            .copy()
        )
        source_rows["power_mw"] = power
        source_rows_output.extend(source_rows.to_dict("records"))

        packages.append(
            {
                "package_id": f"pnnl_v2024_lfp_2023_point_{power:g}mw_4to10h",
                "technology": TECHNOLOGY,
                "source_year": SOURCE_YEAR,
                "estimate_type": ESTIMATE_TYPE,
                "power_scale_bracket_mw": power,
                "duration_window_hr": "4,6,8,10",
                "currency": "USD",
                "currency_base_year": 2023,
                "capex_C_E_usd_per_kwh": capex_C_E,
                "capex_C_P_usd_per_kw": capex_C_P,
                "capex_r2": capex_r2,
                "capex_max_abs_pct_error": capex_max_err,
                "fom_E_usd_per_kwh_year": fom_E,
                "fom_P_usd_per_kw_year": fom_P,
                "fom_r2": fom_r2,
                "fom_max_abs_pct_error": fom_max_err,
                "crep_basis": "DC Storage Block ($/kWh)",
                "crep_reduction_rule": (
                    "arithmetic mean of reported DC Storage Block $/kWh "
                    "at 4,6,8,10 h within same power-scale bracket"
                ),
                "crep_usd_per_kwh": crep_mean,
                "crep_window_median_usd_per_kwh": crep_median,
                "crep_window_min_usd_per_kwh": crep_min,
                "crep_window_max_usd_per_kwh": crep_max,
                "crep_window_range_pct_of_mean": crep_range_pct_of_mean,
                "dc_sb_affine_energy_component_usd_per_kwh_diagnostic": (
                    sb_energy_slope
                ),
                "dc_sb_affine_power_component_usd_per_kw_diagnostic": (
                    sb_power_intercept
                ),
                "dc_sb_affine_r2_diagnostic": sb_r2,
                "dc_sb_affine_max_abs_pct_error_diagnostic": sb_max_err,
                "discount_rate": DISCOUNT_RATE,
                "analysis_horizon_years": ANALYSIS_HORIZON_YR,
                "crf_diagnostic_not_applied": derived_crf,
                "annualized": False,
                "fx_normalized": False,
                "mainline_selected": False,
                "workbook_sha256": workbook_hash,
                "worksheet": "Database",
                "status": "USD2023_CANDIDATE_PACKAGE_FROZEN_BRACKET_NOT_SELECTED",
            }
        )

        print()
        print(f"{power:g} MW candidate package")
        print(
            f"- CAPEX: C_E={capex_C_E:.6f} USD/kWh, "
            f"C_P={capex_C_P:.6f} USD/kW"
        )
        print(
            f"  fit: R2={capex_r2:.8f}, "
            f"max_err={capex_max_err:.4f}%"
        )
        print(
            f"- FOM: FOM_E={fom_E:.6f} USD/kWh-yr, "
            f"FOM_P={fom_P:.6f} USD/kW-yr"
        )
        print(
            f"  fit: R2={fom_r2:.8f}, "
            f"max_err={fom_max_err:.4f}%"
        )
        print(
            f"- C_rep basis: DC Storage Block | "
            f"mean={crep_mean:.4f} USD/kWh"
        )
        print(
            f"  source range={crep_min:.2f} to {crep_max:.2f} USD/kWh "
            f"({crep_range_pct_of_mean:.4f}% of mean)"
        )
        print(
            f"- DC-SB affine diagnostic only: "
            f"E-component={sb_energy_slope:.6f} USD/kWh, "
            f"P-component={sb_power_intercept:.6f} USD/kW, "
            f"R2={sb_r2:.8f}"
        )

    package_df = pd.DataFrame(packages)
    source_df = pd.DataFrame(source_rows_output)

    package_path = (
        output_dir
        / "pnnl_v2024_lfp_usd2023_candidate_cost_packages.csv"
    )
    package_df.to_csv(
        package_path,
        index=False,
        encoding="utf-8-sig",
    )

    source_path = (
        output_dir
        / "pnnl_v2024_lfp_usd2023_candidate_package_source_rows.csv"
    )
    source_df.to_csv(
        source_path,
        index=False,
        encoding="utf-8-sig",
    )

    audit_path = (
        output_dir
        / "pnnl_v2024_lfp_usd2023_candidate_package_audit.txt"
    )

    audit_lines = [
        f"Script version: {SCRIPT_VERSION}",
        "NTUST v7.1 PNNL v2024 LFP USD-2023 Candidate Cost-Package Freeze",
        "",
        f"Workbook: {workbook_path}",
        f"SHA-256: {workbook_hash}",
        "Worksheet: Database",
        "Technology: Lithium-ion LFP",
        "Source year: 2023",
        "Estimate type: Point",
        "Duration window: 4, 6, 8, 10 h",
        "Power-scale brackets: 1 MW, 10 MW",
        "",
        "Frozen source-definition rules:",
        "- CAPEX is represented as C_E*E_nameplate + C_P*P_rating.",
        "- CAPEX coefficients are thesis-specific affine fits of PNNL Total Installed Cost ($/kW) vs duration.",
        "- FOM is represented as FOM_E*E_nameplate + FOM_P*P_rating.",
        "- FOM coefficients are thesis-specific affine fits of PNNL Fixed O&M ($/kW-year) vs duration.",
        "- C_rep component basis is PNNL DC Storage Block ($/kWh).",
        "- C_rep numerical reduction within a bracket is the arithmetic mean across 4/6/8/10 h.",
        "- The same duration window is used for CAPEX, FOM, and C_rep source reduction.",
        "- 1 MW and 10 MW remain separate cost-scale brackets.",
        "- No mainline bracket is selected by this script.",
        "- No interpolation between 1 MW and 10 MW is performed.",
        "- No claim is made that either bracket is an NTUST turnkey quotation.",
        "",
        "Accounting guardrails:",
        "- All values in this file remain raw USD-2023 values.",
        "- No USD->NTD conversion is applied.",
        "- CRF is calculated as a diagnostic but is not applied.",
        "- Annual FOM is already annual and must not receive CRF later.",
        "- C_rep is a cycling-wear economic basis, not a second full replacement cash-flow stream.",
        "- Recycling/decommissioning cost is not included in C_rep.",
        "",
        f"Discount rate: {DISCOUNT_RATE}",
        f"Analysis horizon: {ANALYSIS_HORIZON_YR} years",
        f"Derived CRF diagnostic: {derived_crf:.10f}",
        "",
        "Gate interpretation:",
        "- CAPEX source-definition / linearization: FROZEN FOR BOTH BRACKETS",
        "- FOM source-definition / linearization: FROZEN FOR BOTH BRACKETS",
        "- C_rep component basis: FROZEN AS DC STORAGE BLOCK",
        "- C_rep bracket values: FROZEN AS 4-10h WINDOW MEANS IN USD-2023",
        "- Mainline 1 MW vs 10 MW bracket: NOT SELECTED",
        "- FX normalization: PENDING",
        "- NTD parameter registry population: PENDING",
        "- lambda_k generation: PENDING FX + final bracket selection/usage role",
    ]

    audit_path.write_text(
        "\n".join(audit_lines),
        encoding="utf-8",
    )

    print()
    print("Outputs")
    print(f"- Candidate packages: {package_path}")
    print(f"- Exact source rows: {source_path}")
    print(f"- Audit summary: {audit_path}")

    print()
    print("Gate interpretation")
    print("- CAPEX source-definition: FROZEN")
    print("- FOM source-definition: FROZEN")
    print("- C_rep component basis: FROZEN AS DC STORAGE BLOCK")
    print("- C_rep bracket values: FROZEN IN RAW USD-2023")
    print("- Mainline 1 MW vs 10 MW bracket: NOT SELECTED")
    print("- FX normalization: PENDING")
    print("- lambda_k generation: PENDING")
    print("- Next step: audit and freeze USD->NTD normalization source/rule.")


if __name__ == "__main__":
    main()
