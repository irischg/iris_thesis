from __future__ import annotations

import hashlib
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


SCRIPT_VERSION = "v7.1-pnnl-v2024-lfp-source-provenance-2026-08-20"

EXPECTED_SHA256 = (
    "ef50e87012361e2a10d3df97cb61b93d38eae9255539ca76061e0c5feef501e9"
)

TECHNOLOGY = "Lithium-ion LFP"
SOURCE_YEAR = 2023
ESTIMATE_TYPE = "Point"
POWER_SCALES_MW = {1.0, 10.0}


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def main() -> None:
    print(f"Script version: {SCRIPT_VERSION}")
    print("NTUST v7.1 PNNL v2024 LFP Exact Source-Provenance Audit")
    print(f"Script version: {SCRIPT_VERSION}")
    print()

    root = Path(__file__).resolve().parents[1]

    workbook_path = (
        root / "data" / "reference" / "ESGC_Cost_Performance_Database_v2024.xlsx"
    )

    output_dir = root / "results" / "data_audit"
    output_dir.mkdir(parents=True, exist_ok=True)

    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    workbook_hash = sha256_file(workbook_path)

    print("Workbook provenance")
    print(f"- Workbook: {workbook_path}")
    print(f"- SHA-256: {workbook_hash}")

    if workbook_hash != EXPECTED_SHA256:
        raise RuntimeError(
            "Workbook SHA-256 does not match the audited 11b workbook.\n"
            f"Expected: {EXPECTED_SHA256}\n"
            f"Observed: {workbook_hash}"
        )

    wb = load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
    )

    if "Database" not in wb.sheetnames:
        raise RuntimeError("Worksheet 'Database' not found.")

    ws = wb["Database"]

    expected_headers = [
        "Technology",
        "Year",
        "Power_MW",
        "Duration_hr",
        "Estimate_type",
        "Parameter_category",
        "Parameter",
        "Value",
    ]

    observed_headers = [
        ws.cell(row=1, column=i).value
        for i in range(1, 9)
    ]

    if observed_headers != expected_headers:
        raise RuntimeError(
            "Database headers do not match the audited 11b schema.\n"
            f"Expected: {expected_headers}\n"
            f"Observed: {observed_headers}"
        )

    records: list[dict] = []

    for excel_row, row in enumerate(
        ws.iter_rows(min_row=2, max_col=8, values_only=True),
        start=2,
    ):
        (
            technology,
            year,
            power_mw,
            duration_hr,
            estimate_type,
            parameter_category,
            parameter,
            value,
        ) = row

        if technology != TECHNOLOGY:
            continue

        try:
            year_num = int(year)
        except (TypeError, ValueError):
            continue

        try:
            power_num = float(power_mw)
        except (TypeError, ValueError):
            continue

        if year_num != SOURCE_YEAR:
            continue

        if estimate_type != ESTIMATE_TYPE:
            continue

        if power_num not in POWER_SCALES_MW:
            continue

        try:
            duration_num = float(duration_hr)
        except (TypeError, ValueError):
            duration_num = None

        records.append(
            {
                "excel_row": excel_row,
                "technology": technology,
                "year": year_num,
                "power_mw": power_num,
                "duration_hr": duration_num,
                "estimate_type": estimate_type,
                "parameter_category": parameter_category,
                "parameter": parameter,
                "value": value,
                "source_parameter_cell": f"G{excel_row}",
                "source_value_cell": f"H{excel_row}",
                "worksheet": "Database",
                "workbook_sha256": workbook_hash,
            }
        )

    df = pd.DataFrame(records)

    if df.empty:
        raise RuntimeError(
            "No 2023 Point-estimate LFP rows found for 1 MW / 10 MW."
        )

    df = df.sort_values(
        [
            "power_mw",
            "duration_hr",
            "parameter_category",
            "parameter",
            "excel_row",
        ],
        na_position="last",
    ).reset_index(drop=True)

    all_path = (
        output_dir
        / "pnnl_v2024_lfp_2023_point_1mw_10mw_source_rows.csv"
    )
    df.to_csv(all_path, index=False, encoding="utf-8-sig")

    capital = df[
        df["parameter_category"]
        .astype(str)
        .str.casefold()
        .eq("capital cost")
    ].copy()

    capital_path = (
        output_dir
        / "pnnl_v2024_lfp_2023_point_capital_cost_source_rows.csv"
    )
    capital.to_csv(capital_path, index=False, encoding="utf-8-sig")

    fom_mask = (
        df["parameter"].astype(str).str.contains(
            r"fixed\s*o&m",
            case=False,
            regex=True,
            na=False,
        )
    )
    fom = df[fom_mask].copy()

    fom_path = (
        output_dir
        / "pnnl_v2024_lfp_2023_point_fom_source_rows.csv"
    )
    fom.to_csv(fom_path, index=False, encoding="utf-8-sig")

    crep_mask = (
        df["parameter"].astype(str).str.contains(
            r"storage block|replacement|augmentation",
            case=False,
            regex=True,
            na=False,
        )
    )
    crep_candidates = df[crep_mask].copy()

    crep_path = (
        output_dir
        / "pnnl_v2024_lfp_2023_point_crep_candidate_rows.csv"
    )
    crep_candidates.to_csv(
        crep_path,
        index=False,
        encoding="utf-8-sig",
    )

    print()
    print("Selection rule")
    print(f"- Technology: {TECHNOLOGY}")
    print(f"- Year: {SOURCE_YEAR}")
    print(f"- Estimate type: {ESTIMATE_TYPE}")
    print("- Power scales: 1 MW and 10 MW")
    print("- Duration: NOT frozen by this script")

    print()
    print("Available durations after exact filtering")
    for power in sorted(POWER_SCALES_MW):
        durations = sorted(
            df.loc[df["power_mw"] == power, "duration_hr"]
            .dropna()
            .unique()
            .tolist()
        )
        print(f"- {power:g} MW: {durations}")

    print()
    print("Parameter-category counts")
    counts = (
        df.groupby(
            ["power_mw", "parameter_category"],
            dropna=False,
        )
        .size()
        .reset_index(name="rows")
    )

    for _, r in counts.iterrows():
        print(
            f"- {r['power_mw']:g} MW | "
            f"{r['parameter_category']} | "
            f"rows={int(r['rows'])}"
        )

    print()
    print("Capital Cost parameter names")
    for parameter in sorted(
        capital["parameter"].dropna().astype(str).unique()
    ):
        print(f"- {parameter}")

    print()
    print("Fixed O&M candidate rows")
    if fom.empty:
        print("- NONE FOUND")
    else:
        for _, r in fom.iterrows():
            print(
                f"- row={int(r['excel_row'])} | "
                f"{r['power_mw']:g} MW | "
                f"duration={r['duration_hr']:g} h | "
                f"{r['parameter']} | "
                f"value={r['value']}"
            )

    print()
    print("C_rep-related candidate rows")
    if crep_candidates.empty:
        print("- NONE FOUND")
    else:
        for _, r in crep_candidates.iterrows():
            print(
                f"- row={int(r['excel_row'])} | "
                f"{r['power_mw']:g} MW | "
                f"duration={r['duration_hr']:g} h | "
                f"{r['parameter_category']} | "
                f"{r['parameter']} | "
                f"value={r['value']}"
            )

    audit_path = (
        output_dir
        / "pnnl_v2024_lfp_source_provenance_audit.txt"
    )

    audit_lines = [
        f"Script version: {SCRIPT_VERSION}",
        "NTUST v7.1 PNNL v2024 LFP Exact Source-Provenance Audit",
        "",
        f"Workbook: {workbook_path}",
        f"SHA-256: {workbook_hash}",
        "Worksheet: Database",
        "",
        f"Technology: {TECHNOLOGY}",
        f"Year: {SOURCE_YEAR}",
        f"Estimate_type: {ESTIMATE_TYPE}",
        "Power scales: 1 MW, 10 MW",
        "Duration range: NOT YET FROZEN",
        "",
        f"Selected source rows: {len(df)}",
        f"Capital-cost rows: {len(capital)}",
        f"FOM candidate rows: {len(fom)}",
        f"C_rep-related candidate rows: {len(crep_candidates)}",
        "",
        "Gate interpretation:",
        "- Exact source-cell provenance has been extracted.",
        "- 1 MW / 10 MW are treated as PNNL scale brackets only.",
        "- No NTUST BESS size is inferred from these brackets.",
        "- Duration window is intentionally not frozen here.",
        "- C_E, C_P, FOM, and C_rep are not yet final.",
        "- No USD->NTD conversion is performed.",
        "- No CRF annualization is performed.",
    ]

    audit_path.write_text(
        "\n".join(audit_lines),
        encoding="utf-8",
    )

    print()
    print("Outputs")
    print(f"- All selected source rows: {all_path}")
    print(f"- Capital-cost rows: {capital_path}")
    print(f"- FOM candidates: {fom_path}")
    print(f"- C_rep candidates: {crep_path}")
    print(f"- Audit summary: {audit_path}")

    print()
    print("Gate interpretation")
    print("- Exact source-cell provenance: EXTRACTED")
    print("- 1 MW / 10 MW scale brackets: EXPLICIT")
    print("- Duration window: NOT YET FROZEN")
    print("- C_E / C_P linearization: NOT YET PERFORMED")
    print("- FOM: CANDIDATES EXTRACTED, NOT YET FROZEN")
    print("- C_rep: CANDIDATES EXTRACTED, NOT YET FROZEN")
    print("- FX normalization: NOT YET PERFORMED")
    print("- Next step: review exact rows, then freeze the cost package.")


if __name__ == "__main__":
    main()
