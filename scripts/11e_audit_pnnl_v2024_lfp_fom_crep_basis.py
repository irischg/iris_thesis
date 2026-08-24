from __future__ import annotations

import hashlib
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


SCRIPT_VERSION = "v7.1-pnnl-v2024-lfp-fom-crep-audit-2026-08-22"

EXPECTED_SHA256 = (
    "ef50e87012361e2a10d3df97cb61b93d38eae9255539ca76061e0c5feef501e9"
)

TECHNOLOGY = "Lithium-ion LFP"
SOURCE_YEAR = 2023
ESTIMATE_TYPE = "Point"
POWER_SCALES_MW = (1.0, 10.0)

# Script 11d reproduced the v7.1 registry coefficients from this window.
# This script treats 4/6/8/10 h as the audited CAPEX-linearization window,
# but DOES NOT infer that outage duration beta or physical battery duration
# must equal this window.
AUDITED_DURATION_WINDOW_HR = (4.0, 6.0, 8.0, 10.0)

REPLACEMENT_AUGMENTATION_KEYWORDS = (
    "replace",
    "replacement",
    "augment",
    "augmentation",
)

STORAGE_BLOCK_KEYWORDS = (
    "storage block",
    "battery block",
    "cell",
    "module",
)

O_AND_M_KEYWORDS = (
    "o&m",
    "operation",
    "maintenance",
)

END_OF_LIFE_KEYWORDS = (
    "recycl",
    "decommission",
    "end of life",
    "end-of-life",
    "salvage",
    "disposal",
)


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def norm_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip().casefold()


def has_keyword(value: object, keywords: tuple[str, ...]) -> bool:
    text = norm_text(value)
    return any(keyword in text for keyword in keywords)


def fit_affine_duration(
    durations_hr: np.ndarray,
    values_per_kw: np.ndarray,
) -> dict[str, float]:
    """
    Diagnostic affine fit only:

        y [$/kW-year] = FOM_P [$/kW-year]
                       + FOM_E [$/kWh-year] * duration [h]

    Since E_kWh = P_kW * duration_h, this corresponds to:

        annual FOM = FOM_P * P_kW + FOM_E * E_kWh

    This is NOT frozen as the thesis production formulation by this script.
    """
    x = np.asarray(durations_hr, dtype=float)
    y = np.asarray(values_per_kw, dtype=float)

    X = np.column_stack([np.ones_like(x), x])
    beta, *_ = np.linalg.lstsq(X, y, rcond=None)

    intercept = float(beta[0])
    slope = float(beta[1])

    y_hat = X @ beta
    residual = y - y_hat

    ss_res = float(np.sum(residual**2))
    ss_tot = float(np.sum((y - np.mean(y)) ** 2))
    r2 = 1.0 if ss_tot == 0 else 1.0 - ss_res / ss_tot

    rmse = float(np.sqrt(np.mean(residual**2)))
    mae = float(np.mean(np.abs(residual)))

    with np.errstate(divide="ignore", invalid="ignore"):
        pct_error = np.where(np.abs(y) > 0, residual / y * 100.0, np.nan)

    max_abs_pct_error = float(np.nanmax(np.abs(pct_error)))

    return {
        "fom_power_component_usd_per_kw_year": intercept,
        "fom_energy_component_usd_per_kwh_year": slope,
        "r2": r2,
        "rmse_usd_per_kw_year": rmse,
        "mae_usd_per_kw_year": mae,
        "max_abs_pct_error": max_abs_pct_error,
    }


def main() -> None:
    print(f"Script version: {SCRIPT_VERSION}")
    print("NTUST v7.1 PNNL v2024 LFP FOM / C_rep Source-Definition Audit")
    print(f"Script version: {SCRIPT_VERSION}")
    print()

    root = Path(__file__).resolve().parents[1]
    workbook_path = (
        root / "data" / "reference" / "ESGC_Cost_Performance_Database_v2024.xlsx"
    )
    output_dir = root / "results" / "data_audit"
    output_dir.mkdir(parents=True, exist_ok=True)

    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    workbook_hash = sha256_file(workbook_path)

    print("Workbook provenance")
    print(f"- Workbook: {workbook_path}")
    print(f"- SHA-256: {workbook_hash}")

    if workbook_hash != EXPECTED_SHA256:
        raise RuntimeError(
            "Workbook SHA-256 does not match the audited 11b/11c/11d workbook.\n"
            f"Expected: {EXPECTED_SHA256}\n"
            f"Observed: {workbook_hash}"
        )

    wb = load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
    )

    if "Database" not in wb.sheetnames:
        raise RuntimeError("Worksheet 'Database' not found.")

    ws = wb["Database"]

    expected_headers = [
        "Technology",
        "Year",
        "Power_MW",
        "Duration_hr",
        "Estimate_type",
        "Parameter_category",
        "Parameter",
        "Value",
    ]

    observed_headers = [
        ws.cell(row=1, column=i).value
        for i in range(1, 9)
    ]

    if observed_headers != expected_headers:
        raise RuntimeError(
            "Database headers do not match the audited schema.\n"
            f"Expected: {expected_headers}\n"
            f"Observed: {observed_headers}"
        )

    records: list[dict] = []

    for excel_row, row in enumerate(
        ws.iter_rows(min_row=2, max_col=8, values_only=True),
        start=2,
    ):
        (
            technology,
            year,
            power_mw,
            duration_hr,
            estimate_type,
            parameter_category,
            parameter,
            value,
        ) = row

        if technology != TECHNOLOGY:
            continue

        try:
            year_num = int(year)
            power_num = float(power_mw)
            duration_num = float(duration_hr)
        except (TypeError, ValueError):
            continue

        if year_num != SOURCE_YEAR:
            continue
        if estimate_type != ESTIMATE_TYPE:
            continue
        if power_num not in POWER_SCALES_MW:
            continue
        if duration_num not in AUDITED_DURATION_WINDOW_HR:
            continue

        records.append(
            {
                "excel_row": excel_row,
                "technology": technology,
                "year": year_num,
                "power_mw": power_num,
                "duration_hr": duration_num,
                "estimate_type": estimate_type,
                "parameter_category": parameter_category,
                "parameter": parameter,
                "value": value,
                "source_parameter_cell": f"G{excel_row}",
                "source_value_cell": f"H{excel_row}",
                "worksheet": "Database",
                "workbook_sha256": workbook_hash,
            }
        )

    df = pd.DataFrame(records)

    if df.empty:
        raise RuntimeError("No audited-window PNNL LFP rows were extracted.")

    df = df.sort_values(
        [
            "power_mw",
            "duration_hr",
            "parameter_category",
            "parameter",
            "excel_row",
        ],
        na_position="last",
    ).reset_index(drop=True)

    all_rows_path = (
        output_dir
        / "pnnl_v2024_lfp_4to10h_1mw_10mw_all_source_rows.csv"
    )
    df.to_csv(all_rows_path, index=False, encoding="utf-8-sig")

    inventory = (
        df.groupby(["parameter_category", "parameter"], dropna=False)
        .agg(
            n_rows=("excel_row", "size"),
            n_power_scales=("power_mw", "nunique"),
            n_durations=("duration_hr", "nunique"),
            min_value=("value", "min"),
            max_value=("value", "max"),
        )
        .reset_index()
        .sort_values(["parameter_category", "parameter"])
        .reset_index(drop=True)
    )

    inventory_path = (
        output_dir
        / "pnnl_v2024_lfp_4to10h_parameter_inventory.csv"
    )
    inventory.to_csv(inventory_path, index=False, encoding="utf-8-sig")

    text_for_flag = (
        df["parameter_category"].fillna("").astype(str)
        + " | "
        + df["parameter"].fillna("").astype(str)
    )

    explicit_replacement_mask = text_for_flag.map(
        lambda x: has_keyword(x, REPLACEMENT_AUGMENTATION_KEYWORDS)
    )
    storage_block_mask = text_for_flag.map(
        lambda x: has_keyword(x, STORAGE_BLOCK_KEYWORDS)
    )
    om_mask = text_for_flag.map(
        lambda x: has_keyword(x, O_AND_M_KEYWORDS)
    )
    eol_mask = text_for_flag.map(
        lambda x: has_keyword(x, END_OF_LIFE_KEYWORDS)
    )

    explicit_replacement = df[explicit_replacement_mask].copy()
    storage_block = df[storage_block_mask].copy()
    om_candidates = df[om_mask].copy()
    eol_candidates = df[eol_mask].copy()

    explicit_replacement_path = (
        output_dir
        / "pnnl_v2024_lfp_explicit_replacement_augmentation_rows.csv"
    )
    explicit_replacement.to_csv(
        explicit_replacement_path,
        index=False,
        encoding="utf-8-sig",
    )

    storage_block_path = (
        output_dir
        / "pnnl_v2024_lfp_storage_block_candidate_rows.csv"
    )
    storage_block.to_csv(
        storage_block_path,
        index=False,
        encoding="utf-8-sig",
    )

    om_path = (
        output_dir
        / "pnnl_v2024_lfp_om_candidate_rows.csv"
    )
    om_candidates.to_csv(
        om_path,
        index=False,
        encoding="utf-8-sig",
    )

    eol_path = (
        output_dir
        / "pnnl_v2024_lfp_end_of_life_candidate_rows.csv"
    )
    eol_candidates.to_csv(
        eol_path,
        index=False,
        encoding="utf-8-sig",
    )

    fixed_fom = df[
        df["parameter"]
        .fillna("")
        .astype(str)
        .str.casefold()
        .eq("fixed o&m ($/kw-year)")
    ].copy()

    fom_fit_records: list[dict] = []

    for power in POWER_SCALES_MW:
        sub = fixed_fom[fixed_fom["power_mw"] == power].copy()
        observed = tuple(sorted(sub["duration_hr"].tolist()))
        expected = tuple(sorted(AUDITED_DURATION_WINDOW_HR))

        if observed != expected:
            continue

        values = sub["value"].astype(float).to_numpy()

        fit = fit_affine_duration(
            sub["duration_hr"].to_numpy(dtype=float),
            values,
        )

        fom_fit_records.append(
            {
                "power_mw": power,
                "durations_hr": ",".join(
                    f"{d:g}" for d in AUDITED_DURATION_WINDOW_HR
                ),
                **fit,
                "status": "DIAGNOSTIC_ONLY_NOT_FROZEN",
            }
        )

    fom_fits = pd.DataFrame(fom_fit_records)

    fom_fit_path = (
        output_dir
        / "pnnl_v2024_lfp_fom_4to10h_affine_diagnostics.csv"
    )
    fom_fits.to_csv(
        fom_fit_path,
        index=False,
        encoding="utf-8-sig",
    )

    dc_storage_block = df[
        df["parameter"]
        .fillna("")
        .astype(str)
        .str.casefold()
        .eq("dc storage block ($/kwh)")
    ].copy()

    crep_diag_records: list[dict] = []

    for power in POWER_SCALES_MW:
        sub = dc_storage_block[
            dc_storage_block["power_mw"] == power
        ].copy()

        observed = tuple(sorted(sub["duration_hr"].tolist()))
        expected = tuple(sorted(AUDITED_DURATION_WINDOW_HR))

        if observed != expected:
            continue

        values = sub["value"].astype(float)

        crep_diag_records.append(
            {
                "power_mw": power,
                "durations_hr": ",".join(
                    f"{d:g}" for d in AUDITED_DURATION_WINDOW_HR
                ),
                "min_usd_per_kwh": float(values.min()),
                "max_usd_per_kwh": float(values.max()),
                "mean_usd_per_kwh": float(values.mean()),
                "median_usd_per_kwh": float(values.median()),
                "duration_4h_usd_per_kwh": float(
                    sub.loc[sub["duration_hr"] == 4.0, "value"].iloc[0]
                ),
                "duration_6h_usd_per_kwh": float(
                    sub.loc[sub["duration_hr"] == 6.0, "value"].iloc[0]
                ),
                "duration_8h_usd_per_kwh": float(
                    sub.loc[sub["duration_hr"] == 8.0, "value"].iloc[0]
                ),
                "duration_10h_usd_per_kwh": float(
                    sub.loc[sub["duration_hr"] == 10.0, "value"].iloc[0]
                ),
                "status": "CANDIDATE_BASIS_ONLY_NOT_FROZEN",
            }
        )

    crep_diag = pd.DataFrame(crep_diag_records)

    crep_diag_path = (
        output_dir
        / "pnnl_v2024_lfp_dc_storage_block_crep_diagnostics.csv"
    )
    crep_diag.to_csv(
        crep_diag_path,
        index=False,
        encoding="utf-8-sig",
    )

    print()
    print("Audit scope")
    print(f"- Technology: {TECHNOLOGY}")
    print(f"- Year: {SOURCE_YEAR}")
    print(f"- Estimate type: {ESTIMATE_TYPE}")
    print("- Power scales: 1 MW and 10 MW")
    print("- Audited cost-linearization window: 4, 6, 8, 10 h")
    print("- This script does NOT select the mainline 1 MW vs 10 MW bracket.")

    print()
    print("Full parameter inventory in audited window")
    for category, group in inventory.groupby(
        "parameter_category",
        dropna=False,
    ):
        print(f"- [{category}]")
        for _, row in group.iterrows():
            print(
                f"  - {row['parameter']} | "
                f"rows={int(row['n_rows'])} | "
                f"min={row['min_value']} | max={row['max_value']}"
            )

    print()
    print("Explicit replacement / augmentation parameter rows")
    if explicit_replacement.empty:
        print("- NONE FOUND in 2023 Point / 1 MW & 10 MW / 4-10 h subset")
    else:
        for _, row in explicit_replacement.iterrows():
            print(
                f"- row={int(row['excel_row'])} | "
                f"{row['power_mw']:g} MW | "
                f"{row['duration_hr']:g} h | "
                f"{row['parameter_category']} | "
                f"{row['parameter']} | value={row['value']}"
            )

    print()
    print("Storage-block-related rows")
    if storage_block.empty:
        print("- NONE FOUND")
    else:
        unique_params = sorted(
            storage_block["parameter"].dropna().astype(str).unique()
        )
        for p in unique_params:
            print(f"- {p}")

    print()
    print("O&M-related parameter rows")
    if om_candidates.empty:
        print("- NONE FOUND")
    else:
        unique_params = sorted(
            om_candidates["parameter"].dropna().astype(str).unique()
        )
        for p in unique_params:
            print(f"- {p}")

    print()
    print("End-of-life-related parameter rows")
    if eol_candidates.empty:
        print("- NONE FOUND")
    else:
        unique_pairs = (
            eol_candidates[
                ["parameter_category", "parameter"]
            ]
            .drop_duplicates()
            .sort_values(["parameter_category", "parameter"])
        )
        for _, row in unique_pairs.iterrows():
            print(
                f"- {row['parameter_category']} | {row['parameter']}"
            )

    print()
    print("Fixed O&M 4-10 h affine diagnostics")
    if fom_fits.empty:
        print("- Unable to fit exact Fixed O&M rows.")
    else:
        for _, row in fom_fits.iterrows():
            print(
                f"- {row['power_mw']:g} MW | "
                f"FOM_E={row['fom_energy_component_usd_per_kwh_year']:.6f} "
                f"USD/kWh-yr | "
                f"FOM_P={row['fom_power_component_usd_per_kw_year']:.6f} "
                f"USD/kW-yr | "
                f"R2={row['r2']:.8f} | "
                f"max_err={row['max_abs_pct_error']:.4f}%"
            )

    print()
    print("DC Storage Block candidate C_rep diagnostics")
    if crep_diag.empty:
        print("- Exact DC Storage Block rows not available.")
    else:
        for _, row in crep_diag.iterrows():
            print(
                f"- {row['power_mw']:g} MW | "
                f"4h={row['duration_4h_usd_per_kwh']:.2f} | "
                f"6h={row['duration_6h_usd_per_kwh']:.2f} | "
                f"8h={row['duration_8h_usd_per_kwh']:.2f} | "
                f"10h={row['duration_10h_usd_per_kwh']:.2f} USD/kWh | "
                f"mean={row['mean_usd_per_kwh']:.4f}"
            )

    explicit_status = (
        "EXPLICIT_ROWS_FOUND_REVIEW_REQUIRED"
        if not explicit_replacement.empty
        else "NO_EXPLICIT_ROWS_FOUND"
    )

    audit_path = (
        output_dir
        / "pnnl_v2024_lfp_fom_crep_source_definition_audit.txt"
    )

    audit_lines = [
        f"Script version: {SCRIPT_VERSION}",
        "NTUST v7.1 PNNL v2024 LFP FOM / C_rep Source-Definition Audit",
        "",
        f"Workbook: {workbook_path}",
        f"SHA-256: {workbook_hash}",
        "Worksheet: Database",
        f"Technology: {TECHNOLOGY}",
        f"Year: {SOURCE_YEAR}",
        f"Estimate type: {ESTIMATE_TYPE}",
        "Power scales: 1 MW, 10 MW",
        "Audited duration window: 4, 6, 8, 10 h",
        "",
        f"Selected rows in audited subset: {len(df)}",
        f"Unique parameter/category combinations: {len(inventory)}",
        (
            "Explicit replacement/augmentation rows: "
            f"{len(explicit_replacement)}"
        ),
        f"Storage-block-related rows: {len(storage_block)}",
        f"O&M-related rows: {len(om_candidates)}",
        f"End-of-life-related rows: {len(eol_candidates)}",
        "",
        f"Explicit replacement/augmentation status: {explicit_status}",
        "",
        "Interpretation guardrails:",
        "- PNNL v2024 is the raw CAPEX/FOM source-of-truth.",
        "- This script does not claim that DC Storage Block is automatically C_rep.",
        "- This script does not claim that replacement/augmentation cash flows belong in the objective.",
        "- B2 remains annualized CAPEX + annual FOM + operating costs including cycling wear.",
        "- No second full replacement/augmentation cash-flow stream is introduced.",
        "- FOM affine decomposition is a diagnostic only until explicitly frozen.",
        "- DC Storage Block statistics are candidate C_rep diagnostics only.",
        "- 1 MW / 10 MW remain cost-scale brackets; mainline bracket is not selected here.",
        "- No USD->NTD conversion is performed.",
        "- No CRF annualization is performed.",
        "- No degradation lambda_k values are generated.",
        "",
        "Next gate:",
        "- Review explicit replacement/augmentation parameters, storage-block rows, and FOM diagnostics.",
        "- Then freeze the FOM representation and C_rep economic basis.",
        "- Only after that should FX normalization / final parameter registry / lambda_k generation proceed.",
    ]

    audit_path.write_text(
        "\n".join(audit_lines),
        encoding="utf-8",
    )

    print()
    print("Outputs")
    print(f"- All audited-window rows: {all_rows_path}")
    print(f"- Parameter inventory: {inventory_path}")
    print(f"- Explicit replacement/augmentation rows: {explicit_replacement_path}")
    print(f"- Storage-block candidates: {storage_block_path}")
    print(f"- O&M candidates: {om_path}")
    print(f"- End-of-life candidates: {eol_path}")
    print(f"- FOM affine diagnostics: {fom_fit_path}")
    print(f"- DC Storage Block C_rep diagnostics: {crep_diag_path}")
    print(f"- Audit summary: {audit_path}")

    print()
    print("Gate interpretation")
    print(f"- Explicit replacement/augmentation scan: {explicit_status}")
    print("- FOM source rows: EXTRACTED")
    print("- FOM affine form: DIAGNOSTIC ONLY, NOT FROZEN")
    print("- DC Storage Block source rows: EXTRACTED")
    print("- C_rep basis: NOT YET FROZEN")
    print("- Mainline 1 MW vs 10 MW bracket: NOT SELECTED")
    print("- FX normalization: NOT YET PERFORMED")
    print("- lambda_k generation: NOT YET PERFORMED")
    print("- Next step: review this output and freeze FOM / C_rep definitions.")


if __name__ == "__main__":
    main()
