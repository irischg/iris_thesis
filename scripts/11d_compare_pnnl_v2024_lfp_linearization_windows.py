from __future__ import annotations

import hashlib
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


SCRIPT_VERSION = "v7.1-pnnl-v2024-lfp-linearization-window-audit-2026-08-20"

EXPECTED_SHA256 = (
    "ef50e87012361e2a10d3df97cb61b93d38eae9255539ca76061e0c5feef501e9"
)

TECHNOLOGY = "Lithium-ion LFP"
SOURCE_YEAR = 2023
ESTIMATE_TYPE = "Point"
POWER_SCALES_MW = (1.0, 10.0)

REQUIRED_PARAMETERS = (
    "Total Installed Cost ($)",
    "Total Installed Cost ($/kW)",
    "Total Installed Cost ($/kWh)",
    "Fixed O&M ($/kW-year)",
    "DC Storage Block ($/kWh)",
)

WINDOWS = {
    "2_to_10h": (2.0, 4.0, 6.0, 8.0, 10.0),
    "4_to_10h": (4.0, 6.0, 8.0, 10.0),
    "2_to_24h": (2.0, 4.0, 6.0, 8.0, 10.0, 24.0),
    "4_to_24h": (4.0, 6.0, 8.0, 10.0, 24.0),
    "all_2_to_100h": (2.0, 4.0, 6.0, 8.0, 10.0, 24.0, 100.0),
}


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def fit_affine_duration(
    durations_hr: np.ndarray,
    y_per_kw: np.ndarray,
) -> dict[str, float]:
    """
    Fit:
        y [$/kW] = C_P [$/kW] + C_E [$/kWh] * duration [h]

    Since E_kWh = P_kW * duration_h:
        total_cost = C_P * P_kW + C_E * E_kWh
    """
    x = np.asarray(durations_hr, dtype=float)
    y = np.asarray(y_per_kw, dtype=float)

    X = np.column_stack([np.ones_like(x), x])
    beta, *_ = np.linalg.lstsq(X, y, rcond=None)

    c_p = float(beta[0])
    c_e = float(beta[1])

    y_hat = X @ beta
    residual = y - y_hat

    ss_res = float(np.sum(residual**2))
    ss_tot = float(np.sum((y - np.mean(y)) ** 2))
    r2 = 1.0 if ss_tot == 0 else 1.0 - ss_res / ss_tot

    rmse = float(np.sqrt(np.mean(residual**2)))
    mae = float(np.mean(np.abs(residual)))
    max_abs_error = float(np.max(np.abs(residual)))

    with np.errstate(divide="ignore", invalid="ignore"):
        pct_error = np.where(np.abs(y) > 0, residual / y * 100.0, np.nan)

    max_abs_pct_error = float(np.nanmax(np.abs(pct_error)))

    return {
        "intercept_c_p_per_kw": c_p,
        "slope_c_e_per_kwh": c_e,
        "r2": r2,
        "rmse_per_kw": rmse,
        "mae_per_kw": mae,
        "max_abs_error_per_kw": max_abs_error,
        "max_abs_pct_error": max_abs_pct_error,
    }


def main() -> None:
    print(f"Script version: {SCRIPT_VERSION}")
    print("NTUST v7.1 PNNL v2024 LFP Linearization-Window Audit")
    print(f"Script version: {SCRIPT_VERSION}")
    print()

    root = Path(__file__).resolve().parents[1]
    workbook_path = (
        root / "data" / "reference" / "ESGC_Cost_Performance_Database_v2024.xlsx"
    )
    output_dir = root / "results" / "data_audit"
    output_dir.mkdir(parents=True, exist_ok=True)

    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    workbook_hash = sha256_file(workbook_path)

    print("Workbook provenance")
    print(f"- Workbook: {workbook_path}")
    print(f"- SHA-256: {workbook_hash}")

    if workbook_hash != EXPECTED_SHA256:
        raise RuntimeError(
            "Workbook SHA-256 does not match the audited 11b/11c workbook.\n"
            f"Expected: {EXPECTED_SHA256}\n"
            f"Observed: {workbook_hash}"
        )

    wb = load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
    )

    if "Database" not in wb.sheetnames:
        raise RuntimeError("Worksheet 'Database' not found.")

    ws = wb["Database"]

    expected_headers = [
        "Technology",
        "Year",
        "Power_MW",
        "Duration_hr",
        "Estimate_type",
        "Parameter_category",
        "Parameter",
        "Value",
    ]

    observed_headers = [
        ws.cell(row=1, column=i).value
        for i in range(1, 9)
    ]

    if observed_headers != expected_headers:
        raise RuntimeError(
            "Database headers do not match the audited schema.\n"
            f"Expected: {expected_headers}\n"
            f"Observed: {observed_headers}"
        )

    records: list[dict] = []

    for excel_row, row in enumerate(
        ws.iter_rows(min_row=2, max_col=8, values_only=True),
        start=2,
    ):
        (
            technology,
            year,
            power_mw,
            duration_hr,
            estimate_type,
            parameter_category,
            parameter,
            value,
        ) = row

        if technology != TECHNOLOGY:
            continue

        try:
            year_num = int(year)
            power_num = float(power_mw)
            duration_num = float(duration_hr)
            value_num = float(value)
        except (TypeError, ValueError):
            continue

        if year_num != SOURCE_YEAR:
            continue
        if estimate_type != ESTIMATE_TYPE:
            continue
        if power_num not in POWER_SCALES_MW:
            continue
        if parameter not in REQUIRED_PARAMETERS:
            continue

        records.append(
            {
                "excel_row": excel_row,
                "power_mw": power_num,
                "duration_hr": duration_num,
                "parameter_category": parameter_category,
                "parameter": parameter,
                "value": value_num,
                "source_value_cell": f"H{excel_row}",
            }
        )

    raw = pd.DataFrame(records)

    if raw.empty:
        raise RuntimeError("No required PNNL rows were extracted.")

    duplicate_check = (
        raw.groupby(["power_mw", "duration_hr", "parameter"])
        .size()
        .reset_index(name="n")
    )
    bad_dupes = duplicate_check[duplicate_check["n"] != 1]
    if not bad_dupes.empty:
        raise RuntimeError(
            "Expected exactly one row per power/duration/parameter, but found:\n"
            + bad_dupes.to_string(index=False)
        )

    pivot = (
        raw.pivot(
            index=["power_mw", "duration_hr"],
            columns="parameter",
            values="value",
        )
        .reset_index()
        .sort_values(["power_mw", "duration_hr"])
        .reset_index(drop=True)
    )

    source_rows = (
        raw.pivot(
            index=["power_mw", "duration_hr"],
            columns="parameter",
            values="excel_row",
        )
        .reset_index()
        .sort_values(["power_mw", "duration_hr"])
        .reset_index(drop=True)
    )

    missing_parameters = [
        p for p in REQUIRED_PARAMETERS if p not in pivot.columns
    ]
    if missing_parameters:
        raise RuntimeError(
            f"Required parameters missing after pivot: {missing_parameters}"
        )

    pivot["power_kw"] = pivot["power_mw"] * 1000.0
    pivot["energy_kwh"] = pivot["power_kw"] * pivot["duration_hr"]

    pivot["tic_per_kw_from_total_dollars"] = (
        pivot["Total Installed Cost ($)"] / pivot["power_kw"]
    )
    pivot["tic_per_kw_from_per_kwh"] = (
        pivot["Total Installed Cost ($/kWh)"] * pivot["duration_hr"]
    )

    pivot["tic_per_kw_diff_total_vs_reported"] = (
        pivot["tic_per_kw_from_total_dollars"]
        - pivot["Total Installed Cost ($/kW)"]
    )
    pivot["tic_per_kw_diff_perkwh_vs_reported"] = (
        pivot["tic_per_kw_from_per_kwh"]
        - pivot["Total Installed Cost ($/kW)"]
    )

    pivot["tic_total_rel_error_pct"] = (
        pivot["tic_per_kw_diff_total_vs_reported"]
        / pivot["Total Installed Cost ($/kW)"]
        * 100.0
    )
    pivot["tic_perkwh_rel_error_pct"] = (
        pivot["tic_per_kw_diff_perkwh_vs_reported"]
        / pivot["Total Installed Cost ($/kW)"]
        * 100.0
    )

    consistency_path = (
        output_dir
        / "pnnl_v2024_lfp_total_installed_cost_consistency.csv"
    )
    pivot.to_csv(consistency_path, index=False, encoding="utf-8-sig")

    source_rows_path = (
        output_dir
        / "pnnl_v2024_lfp_linearization_source_excel_rows.csv"
    )
    source_rows.to_csv(source_rows_path, index=False, encoding="utf-8-sig")

    print()
    print("Total Installed Cost representation consistency")
    for power in POWER_SCALES_MW:
        sub = pivot[pivot["power_mw"] == power]
        max_total = sub["tic_total_rel_error_pct"].abs().max()
        max_perkwh = sub["tic_perkwh_rel_error_pct"].abs().max()

        print(
            f"- {power:g} MW | "
            f"max |total$ -> $/kW discrepancy| = {max_total:.6f}% | "
            f"max |$/kWh * duration -> $/kW discrepancy| = {max_perkwh:.6f}%"
        )

    fit_records: list[dict] = []
    prediction_records: list[dict] = []
    crep_records: list[dict] = []

    for power in POWER_SCALES_MW:
        power_df = pivot[pivot["power_mw"] == power].copy()

        print()
        print(f"{power:g} MW candidate linearizations")

        for window_name, durations in WINDOWS.items():
            sub = power_df[
                power_df["duration_hr"].isin(durations)
            ].copy()

            observed_duration_set = tuple(
                sorted(sub["duration_hr"].tolist())
            )
            expected_duration_set = tuple(sorted(durations))

            if observed_duration_set != expected_duration_set:
                print(
                    f"- {window_name}: SKIPPED; expected={expected_duration_set}, "
                    f"observed={observed_duration_set}"
                )
                continue

            capex_fit = fit_affine_duration(
                sub["duration_hr"].to_numpy(),
                sub["Total Installed Cost ($/kW)"].to_numpy(),
            )

            fom_fit = fit_affine_duration(
                sub["duration_hr"].to_numpy(),
                sub["Fixed O&M ($/kW-year)"].to_numpy(),
            )

            dc_block = sub["DC Storage Block ($/kWh)"]

            record = {
                "power_mw": power,
                "window": window_name,
                "durations_hr": ",".join(
                    f"{d:g}" for d in expected_duration_set
                ),
                "n_points": len(sub),
                "capex_C_P_usd_per_kw": capex_fit["intercept_c_p_per_kw"],
                "capex_C_E_usd_per_kwh": capex_fit["slope_c_e_per_kwh"],
                "capex_r2": capex_fit["r2"],
                "capex_rmse_usd_per_kw": capex_fit["rmse_per_kw"],
                "capex_mae_usd_per_kw": capex_fit["mae_per_kw"],
                "capex_max_abs_error_usd_per_kw": capex_fit[
                    "max_abs_error_per_kw"
                ],
                "capex_max_abs_pct_error": capex_fit[
                    "max_abs_pct_error"
                ],
                "fom_power_component_usd_per_kw_year": fom_fit[
                    "intercept_c_p_per_kw"
                ],
                "fom_energy_component_usd_per_kwh_year": fom_fit[
                    "slope_c_e_per_kwh"
                ],
                "fom_r2": fom_fit["r2"],
                "fom_rmse_usd_per_kw_year": fom_fit["rmse_per_kw"],
                "fom_max_abs_pct_error": fom_fit["max_abs_pct_error"],
                "dc_storage_block_min_usd_per_kwh": float(dc_block.min()),
                "dc_storage_block_max_usd_per_kwh": float(dc_block.max()),
                "dc_storage_block_mean_usd_per_kwh": float(dc_block.mean()),
                "dc_storage_block_median_usd_per_kwh": float(dc_block.median()),
            }
            fit_records.append(record)

            print(
                f"- {window_name}: "
                f"C_E={record['capex_C_E_usd_per_kwh']:.6f} USD/kWh, "
                f"C_P={record['capex_C_P_usd_per_kw']:.6f} USD/kW, "
                f"R2={record['capex_r2']:.8f}, "
                f"RMSE={record['capex_rmse_usd_per_kw']:.6f} USD/kW, "
                f"max_err={record['capex_max_abs_pct_error']:.4f}%"
            )
            print(
                f"  FOM diagnostic: "
                f"energy={record['fom_energy_component_usd_per_kwh_year']:.6f} "
                f"USD/kWh-yr, "
                f"power={record['fom_power_component_usd_per_kw_year']:.6f} "
                f"USD/kW-yr, "
                f"R2={record['fom_r2']:.8f}"
            )
            print(
                f"  DC Storage Block range: "
                f"{record['dc_storage_block_min_usd_per_kwh']:.2f} to "
                f"{record['dc_storage_block_max_usd_per_kwh']:.2f} USD/kWh"
            )

            c_e = record["capex_C_E_usd_per_kwh"]
            c_p = record["capex_C_P_usd_per_kw"]

            fom_e = record["fom_energy_component_usd_per_kwh_year"]
            fom_p = record["fom_power_component_usd_per_kw_year"]

            for _, row in sub.iterrows():
                d = float(row["duration_hr"])
                actual_capex = float(row["Total Installed Cost ($/kW)"])
                predicted_capex = c_p + c_e * d

                actual_fom = float(row["Fixed O&M ($/kW-year)"])
                predicted_fom = fom_p + fom_e * d

                prediction_records.append(
                    {
                        "power_mw": power,
                        "window": window_name,
                        "duration_hr": d,
                        "actual_tic_usd_per_kw": actual_capex,
                        "predicted_tic_usd_per_kw": predicted_capex,
                        "capex_residual_usd_per_kw": (
                            actual_capex - predicted_capex
                        ),
                        "actual_fom_usd_per_kw_year": actual_fom,
                        "predicted_fom_usd_per_kw_year": predicted_fom,
                        "fom_residual_usd_per_kw_year": (
                            actual_fom - predicted_fom
                        ),
                    }
                )

            crep_records.append(
                {
                    "power_mw": power,
                    "window": window_name,
                    "durations_hr": ",".join(
                        f"{d:g}" for d in expected_duration_set
                    ),
                    "dc_storage_block_min_usd_per_kwh": float(dc_block.min()),
                    "dc_storage_block_max_usd_per_kwh": float(dc_block.max()),
                    "dc_storage_block_mean_usd_per_kwh": float(dc_block.mean()),
                    "dc_storage_block_median_usd_per_kwh": float(dc_block.median()),
                    "note": (
                        "Diagnostic only. This script does not freeze C_rep."
                    ),
                }
            )

    fits = pd.DataFrame(fit_records)
    predictions = pd.DataFrame(prediction_records)
    crep_diag = pd.DataFrame(crep_records)

    fit_path = (
        output_dir
        / "pnnl_v2024_lfp_linearization_window_comparison.csv"
    )
    fits.to_csv(fit_path, index=False, encoding="utf-8-sig")

    prediction_path = (
        output_dir
        / "pnnl_v2024_lfp_linearization_predictions.csv"
    )
    predictions.to_csv(
        prediction_path,
        index=False,
        encoding="utf-8-sig",
    )

    crep_path = (
        output_dir
        / "pnnl_v2024_lfp_crep_window_diagnostics.csv"
    )
    crep_diag.to_csv(crep_path, index=False, encoding="utf-8-sig")

    print()
    print("Cross-scale comparison by candidate window")
    for window_name in WINDOWS:
        sub = fits[fits["window"] == window_name].copy()
        if len(sub) != len(POWER_SCALES_MW):
            continue

        row_1 = sub[sub["power_mw"] == 1.0].iloc[0]
        row_10 = sub[sub["power_mw"] == 10.0].iloc[0]

        ce_diff_pct = (
            (row_10["capex_C_E_usd_per_kwh"]
             - row_1["capex_C_E_usd_per_kwh"])
            / row_1["capex_C_E_usd_per_kwh"]
            * 100.0
        )
        cp_diff_pct = (
            (row_10["capex_C_P_usd_per_kw"]
             - row_1["capex_C_P_usd_per_kw"])
            / row_1["capex_C_P_usd_per_kw"]
            * 100.0
        )

        print(
            f"- {window_name}: "
            f"C_E 10MW-vs-1MW={ce_diff_pct:+.3f}% | "
            f"C_P 10MW-vs-1MW={cp_diff_pct:+.3f}%"
        )

    audit_path = (
        output_dir
        / "pnnl_v2024_lfp_linearization_window_audit.txt"
    )

    max_total_discrep = float(
        pivot["tic_total_rel_error_pct"].abs().max()
    )
    max_perkwh_discrep = float(
        pivot["tic_perkwh_rel_error_pct"].abs().max()
    )

    audit_lines = [
        f"Script version: {SCRIPT_VERSION}",
        "NTUST v7.1 PNNL v2024 LFP Linearization-Window Audit",
        "",
        f"Workbook: {workbook_path}",
        f"SHA-256: {workbook_hash}",
        "Technology: Lithium-ion LFP",
        "Year: 2023",
        "Estimate type: Point",
        "Power scales: 1 MW, 10 MW",
        "",
        "Model audited:",
        "  Total Installed Cost ($/kW) = C_P ($/kW) + C_E ($/kWh) * duration (h)",
        "Equivalent:",
        "  CAPEX = C_P * P_kW + C_E * E_kWh",
        "",
        "Candidate windows compared:",
    ]

    for name, durations in WINDOWS.items():
        audit_lines.append(
            f"  - {name}: {', '.join(f'{d:g}' for d in durations)} h"
        )

    audit_lines += [
        "",
        "Consistency diagnostics:",
        (
            "  - Max discrepancy, Total Installed Cost ($) / power "
            f"vs reported $/kW: {max_total_discrep:.8f}%"
        ),
        (
            "  - Max discrepancy, reported $/kWh * duration "
            f"vs reported $/kW: {max_perkwh_discrep:.8f}%"
        ),
        "",
        "Gate interpretation:",
        "  - This script compares candidate duration windows; it does NOT freeze one.",
        "  - C_E and C_P remain thesis-derived planning coefficients.",
        "  - 1 MW and 10 MW remain PNNL scale brackets only.",
        "  - FOM energy/power separation is reported as a diagnostic only.",
        "  - C_rep is NOT frozen; DC Storage Block ranges are diagnostics only.",
        "  - No USD->NTD conversion is performed.",
        "  - No CRF annualization is performed.",
        "  - Next step is to review fit quality and choose/freeze the mainline window.",
    ]

    audit_path.write_text(
        "\n".join(audit_lines),
        encoding="utf-8",
    )

    print()
    print("Outputs")
    print(f"- Cost consistency: {consistency_path}")
    print(f"- Source Excel rows: {source_rows_path}")
    print(f"- Window comparison: {fit_path}")
    print(f"- Fit predictions: {prediction_path}")
    print(f"- C_rep diagnostics: {crep_path}")
    print(f"- Audit summary: {audit_path}")

    print()
    print("Gate interpretation")
    print("- Candidate duration windows: COMPARED")
    print("- Mainline duration window: NOT YET FROZEN")
    print("- C_E / C_P: CANDIDATE FITS ONLY")
    print("- FOM: DIAGNOSTIC FITS ONLY")
    print("- C_rep: DIAGNOSTIC RANGE ONLY")
    print("- FX normalization: NOT YET PERFORMED")
    print("- CRF annualization: NOT YET PERFORMED")
    print("- Next step: review fit quality, then freeze the cost package.")


if __name__ == "__main__":
    main()
